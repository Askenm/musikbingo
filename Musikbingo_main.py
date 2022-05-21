import email
import streamlit as st
import streamlit.components.v1 as components
from spotipy.oauth2 import SpotifyClientCredentials
import spotipy
import random
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
import os
from zipfile import ZipFile
import smtplib
from os.path import basename
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.utils import COMMASPACE, formatdate
import smtplib, ssl
import shutil
import credentials as c


def send_mail(send_from, send_to, subject, text, files=None,
              server="127.0.0.2"):
    assert isinstance(send_to, list)

    msg = MIMEMultipart()
    msg['From'] = send_from
    msg['To'] = COMMASPACE.join(send_to)
    msg['Date'] = formatdate(localtime=True)
    msg['Subject'] = subject

    msg.attach(MIMEText(text))

    for f in files or []:
        with open(f, "rb") as fil:
            part = MIMEApplication(
                fil.read(),
                Name=basename(f)
            )
        # After the file is closed
        part['Content-Disposition'] = 'attachment; filename="%s"' % basename(f)
        msg.attach(part)

    context = ssl.create_default_context()
    with smtplib.SMTP('smtp.gmail.com', 587) as server:
        server.ehlo()  # Can be omitted
        server.starttls(context=context)
        server.ehlo()  # Can be omitted
        server.login(send_from,c.mail_pswrd)
        server.sendmail(send_from, send_to, msg.as_string())
        server.close()


def initialize_spotify():
    cid = c.cid
    secret  = c.secret
    client_credentials_manager = SpotifyClientCredentials(client_id=cid, client_secret=secret)
    sp = spotipy.Spotify(client_credentials_manager = client_credentials_manager)

    return sp


def get_tile_values_from_playlist(playlist_link,sp):
    playlist_URI = playlist_link.split("/")[-1].split("?")[0]
    try:
        track_uris = [x["track"]["uri"] for x in sp.playlist_tracks(playlist_URI)["items"]]
    except:
        st.info('Please supply a valid spotify playlist link')
        return [],[]

    value_strings = [track["track"]["name"]+'\n'+track["track"]["artists"][0]["name"]
                    for track in sp.playlist_tracks(playlist_URI)["items"]]


    ixes = [i for i in range(len(value_strings))]

    return value_strings,ixes

def create_bingo_dfs(value_strings,num_plader,max_songs_pr_sheet = 10):
    bingoplade_dfs = []

    for plade in range(num_plader):
        chosen_songs = []
        cur_plade = {'kol_{}'.format(i) : [] for i in range(7)}
        print(plade)
        num_songs_pr_row = {i:3 for i in range(3)}
        num_songs_pr_row[random.choice([0,1,2])]+=1

        for row in range(3):
            blank_ixes = [i for i in range(7)]
            song_ixes = [blank_ixes.pop(random.choice([i for i in range(len(blank_ixes))])) for j in range(num_songs_pr_row[row])]

            for col in range(7):
                if col in song_ixes:
                    song = None
                    while song == None or song in chosen_songs:
                        song = random.choice(value_strings)
                        
                    
                    cur_plade['kol_{}'.format(col)].append(song)
                    chosen_songs.append(song)
                else:
                    cur_plade['kol_{}'.format(col)].append('')


        bingoplade_dfs.append(pd.DataFrame.from_dict(cur_plade))

    return bingoplade_dfs


def save_dfs_to_docx(bingoplade_dfs):
    dir_name = 'bingoplader_'+str(random.random())[2:]
    os.mkdir(dir_name)
    zipObj = ZipFile(dir_name+'/bingoplader.zip', 'w')
    for ix,plade_df in enumerate(bingoplade_dfs):
        for col in plade_df.columns:
            plade_df[plade_df[col][0]] = plade_df[col]
            plade_df.drop(col,axis=1,inplace=True)
        plade_df.drop(0,axis=0,inplace=True)
        document = Document()
        style = document.styles['Normal']
        font = style.font
        font.name = 'Calibri'
        section = document.sections[-1]
        new_width, new_height = section.page_height, section.page_width
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = new_width
        section.page_height = new_height



        t = document.add_table(plade_df.shape[0]+1, plade_df.shape[1])

        # add the header rows.
        for j in range(plade_df.shape[-1]):
            t.cell(0,j).text = plade_df.columns[j]

        # add the rest of the data frame
        for i in range(plade_df.shape[0]):
            for j in range(plade_df.shape[-1]):
                t.cell(i+1,j).text = str(plade_df.values[i,j])

        document.save(dir_name+'/bingoplade_{}.docx'.format(ix+1))
        zipObj.write(dir_name+'/bingoplade_{}.docx'.format(ix+1))
        
    zipObj.close()

    return dir_name


def main():
    st.header('Musikbingo')


    playlist_link = st.text_input('Spotify Playlist Link:')
    email_address = st.text_input('Modtager Email:')
    num_plader = st.slider("antal bingoplader")



    sp = initialize_spotify()
    
    value_strings,ixes = get_tile_values_from_playlist(playlist_link,sp)


    

    # If this is the first run of the app, create 'lav_bingo' in the session_state dict
    if 'lav_bingo' not in st.session_state:
        st.session_state['lav_bingo'] = False

    # Letting the app be activated by the session state rather than a button is much more stable
    if st.button('Lav Bingoplader!'):
        st.session_state['lav_bingo'] = True

    # The main functionality is activated
    if st.session_state['lav_bingo'] and playlist_link and email_address and num_plader:
        

        
        bingoplade_dfs = create_bingo_dfs(value_strings,num_plader)
        
        st.info("Laver Bingoplader fra playlisten {}...".format(sp.playlist(playlist_link)['name']))
        
        dir_name = save_dfs_to_docx(bingoplade_dfs)

        st.info("Sender bingoplader til "+email_address)
        send_mail('aske.osv@gmail.com',[email_address],'Musikbingo','Go Quiz!',[dir_name+'/bingoplader.zip'])

        # DELETING ALL FILES SAVED LOCALLY
        shutil.rmtree(dir_name)

        # RESETTING STATE ST APP DOESNT TRY TO RUN AGAIN WITH SAME VALUES
        st.session_state['lav_bingo'] = False

        # C3L38R4T3
        st.balloons()

    elif not st.session_state['lav_bingo']:
        pass
    elif not playlist_link:
        st.info('Indtast et gyldigt spotify playliste-link')

    elif not email_address:
        st.info('Indtast en gyldig email addresse så du kan modtage bingopladerne')

    elif not num_plader:
        st.info('PVælg et antal bingoplader (større end nul, self)')



if __name__ == '__main__':
    main()